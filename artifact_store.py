"""
artifact_store.py — 费佳生成的 HTML/Markdown/Word 产物的统一存取层。

app.py（预览/下载端点）和 gateway.py（生成工具）都从这里读写，
两边共享同一个 memories.db + artifacts/ 目录，不重复实现。
"""
import sqlite3
import uuid
import os
from pathlib import Path

DB_PATH = '/opt/frontend/memories.db'
STORE_DIR = '/opt/frontend/artifacts'

EXT_BY_TYPE = {'html': 'html', 'markdown': 'md', 'docx': 'docx'}
MIME_BY_TYPE = {
    'html': 'text/html',
    'markdown': 'text/markdown',
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
}


def _artifact_path(filename):
    """Resolve a stored basename without allowing traversal or sibling prefixes."""
    name = str(filename or '')
    if not name or '/' in name or '\\' in name or Path(name).name != name:
        return None
    base = Path(STORE_DIR).resolve()
    candidate = (base / name).resolve()
    try:
        candidate.relative_to(base)
    except ValueError:
        return None
    return candidate


def _init_table():
    os.makedirs(STORE_DIR, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.execute('''CREATE TABLE IF NOT EXISTS artifacts (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        type TEXT NOT NULL,
        title TEXT NOT NULL,
        filename TEXT NOT NULL,
        size INTEGER NOT NULL,
        created_at DATETIME DEFAULT (datetime('now','+8 hours'))
    )''')
    conn.commit()
    conn.close()


_init_table()


def _markdown_to_docx_bytes(title, md_text):
    """把 markdown 文字粗略转成一份 docx：标题/正文/列表，不追求排版还原度。"""
    import docx
    doc = docx.Document()
    doc.add_heading(title, level=1)
    for raw_line in md_text.split('\n'):
        line = raw_line.rstrip()
        if not line.strip():
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.strip().startswith(('- ', '* ')):
            doc.add_paragraph(line.strip()[2:].strip(), style='List Bullet')
        elif line.strip()[:2].isdigit() and '. ' in line[:5]:
            doc.add_paragraph(line.strip().split('. ', 1)[-1], style='List Number')
        else:
            doc.add_paragraph(line)
    import io
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def save(atype, title, content):
    """
    生成并保存一个 artifact。
    atype: 'html' | 'markdown' | 'docx'
    content: html 是完整 HTML 字符串；markdown/docx 是 markdown 格式文字
             （docx 由 markdown 转换而来，费佳写 markdown 就行，不用管 docx 格式细节）
    返回 {'id', 'type', 'title', 'size'}
    """
    if atype not in EXT_BY_TYPE:
        raise ValueError('未知 artifact 类型: ' + str(atype))
    ext = EXT_BY_TYPE[atype]
    fname = 'a_' + uuid.uuid4().hex[:12] + '.' + ext
    fpath = _artifact_path(fname)
    if fpath is None:
        raise ValueError('无效 artifact 文件名')

    if atype == 'docx':
        data = _markdown_to_docx_bytes(title, content)
        with open(fpath, 'wb') as f:
            f.write(data)
        size = len(data)
    else:
        with open(fpath, 'w', encoding='utf-8') as f:
            f.write(content)
        size = len(content.encode('utf-8'))

    conn = sqlite3.connect(DB_PATH)
    cur = conn.execute(
        'INSERT INTO artifacts (type, title, filename, size) VALUES (?,?,?,?)',
        (atype, title, fname, size)
    )
    conn.commit()
    aid = cur.lastrowid
    conn.close()
    return {'id': aid, 'type': atype, 'title': title, 'size': size}


def get(artifact_id):
    """返回 {'id','type','title','filename','size','created_at'} 或 None"""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    row = conn.execute('SELECT * FROM artifacts WHERE id=?', (artifact_id,)).fetchone()
    conn.close()
    return dict(row) if row else None


def read_content(artifact_id):
    """返回 (meta_dict, raw_bytes) 或 (None, None)"""
    meta = get(artifact_id)
    if not meta:
        return None, None
    fpath = _artifact_path(meta['filename'])
    if fpath is None:
        return meta, None
    try:
        with open(fpath, 'rb') as f:
            return meta, f.read()
    except Exception:
        return meta, None


def list_recent(limit=200):
    """Return newest artifacts first. ``limit`` clamped to 1..500."""
    try:
        lim = int(limit)
    except (TypeError, ValueError):
        lim = 200
    lim = max(1, min(lim, 500))
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    rows = conn.execute(
        'SELECT * FROM artifacts ORDER BY datetime(created_at) DESC, id DESC LIMIT ?',
        (lim,),
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def delete(artifact_id, *, allow_stale_metadata=True):
    """Delete one artifact's physical file (if present) and metadata row.

    Returns a status string:
      deleted | not_found | bad_path | unlink_failed | stale_cleared

    ``allow_stale_metadata``: when the DB row exists but the file is already
    gone (or path is unresolvable), still remove the metadata row. Used for
    explicit user deletes and cleaner broken-row cleanup.
    """
    try:
        aid = int(artifact_id)
    except (TypeError, ValueError):
        return 'not_found'
    if aid <= 0:
        return 'not_found'

    meta = get(aid)
    if not meta:
        return 'not_found'

    fpath = _artifact_path(meta.get('filename'))
    if fpath is None:
        if allow_stale_metadata:
            _delete_row(aid)
            return 'stale_cleared'
        return 'bad_path'

    file_existed = fpath.is_file()
    if file_existed:
        try:
            fpath.unlink()
        except OSError:
            return 'unlink_failed'
    elif not allow_stale_metadata:
        return 'not_found'

    _delete_row(aid)
    return 'deleted' if file_existed else 'stale_cleared'


def _delete_row(artifact_id):
    conn = sqlite3.connect(DB_PATH)
    conn.execute('DELETE FROM artifacts WHERE id=?', (int(artifact_id),))
    conn.commit()
    conn.close()
